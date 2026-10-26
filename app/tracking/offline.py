"""
داده‌ی فروش آفلاین (فایل اکسل خود فروشگاه: فاکتورها، تماس‌ها، تنخواه).

فایل در پنل بارگذاری می‌شود و همان لحظه به «آمار جمعی» تبدیل می‌شود:
- برای هر دسته‌ی نقشه: تعداد فاکتور، فروش، سود واقعی، میانگین سود هر فاکتور
- برای هر (دسته، استان): تعداد فاکتور  ← در امتیاز برنامه اثر دارد
- تماس‌ها: نرخ تبدیل تماس به خرید (کل، هر دسته، هر ماه)
- سقف هزینه‌ی مجاز: هر فروش و هر تماس چقدر می‌ارزد

ستون‌های نام، تلفن، آدرس و بانک اصلاً خوانده نمی‌شوند و خود فایل ذخیره نمی‌شود.
"""
import io
import logging
import re
import statistics
from collections import Counter, defaultdict
from datetime import date, datetime, timezone

from app.db.session import get_setting, set_setting

log = logging.getLogger("tracking.offline")
KEY = "offline.summary"
FA = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")

PROVS = ["تهران", "اصفهان", "خراسان رضوی", "فارس", "آذربایجان شرقی", "البرز", "مازندران", "مرکزی", "قزوین", "گیلان",
         "خوزستان", "قم", "آذربایجان غربی", "سمنان", "کرمان", "یزد", "گلستان", "زنجان", "همدان", "کردستان", "اردبیل",
         "لرستان", "کرمانشاه", "هرمزگان", "بوشهر", "ایلام", "خراسان جنوبی", "خراسان شمالی", "سیستان و بلوچستان",
         "چهارمحال و بختیاری", "کهگیلویه و بویراحمد"]
CITY2PROV = {"کرج": "البرز", "ارومیه": "آذربایجان غربی", "مشهد": "خراسان رضوی", "شیراز": "فارس", "رشت": "گیلان",
             "تبریز": "آذربایجان شرقی", "اهواز": "خوزستان", "ساری": "مازندران", "بندرعباس": "هرمزگان", "اراک": "مرکزی",
             "زاهدان": "سیستان و بلوچستان", "سیستان": "سیستان و بلوچستان", "بلوچستان": "سیستان و بلوچستان",
             "بیرجند": "خراسان جنوبی", "بجنورد": "خراسان شمالی", "سنندج": "کردستان", "خرم آباد": "لرستان",
             "شهرکرد": "چهارمحال و بختیاری", "یاسوج": "کهگیلویه و بویراحمد", "گرگان": "گلستان",
             "خراسان": "خراسان رضوی", "آذربایجان": "آذربایجان شرقی", "چهارمحال": "چهارمحال و بختیاری",
             "کهگیلویه": "کهگیلویه و بویراحمد",
             "رباط کریم": "تهران", "شادآباد": "تهران", "اسلامشهر": "تهران", "شهریار": "تهران", "ورامین": "تهران",
             "هشتگرد": "البرز", "نظرآباد": "البرز", "قشم": "هرمزگان", "کیش": "هرمزگان", "مراغه": "آذربایجان شرقی",
             "بهشهر": "مازندران", "آمل": "مازندران", "بابل": "مازندران", "کوهدشت": "لرستان", "بروجرد": "لرستان",
             "دزفول": "خوزستان", "کاشان": "اصفهان", "نجف آباد": "اصفهان", "نیشابور": "خراسان رضوی",
             "رفسنجان": "کرمان", "سیرجان": "کرمان", "زابل": "سیستان و بلوچستان"}

# کلیدواژه ← دسته‌ی نقشه (ترتیب مهم است: خاص‌ترها اول)
KW = [("PNEU", "پنوماتیک|جک |شیر برقی|واحد مراقبت|سنسور جک|سایلنسر|کوپلینگ پنوماتیک|شیلنگ پنوماتیک|کمپرسور|ایرکنترل|مایندمن|سنگ فیلتر"),
      ("HYD_GEAR_VANE", "پمپ دنده|پمپ پره|پمپ هیدرولیک|کارتریج"),
      ("HYD_MOTOR_PISTON", "هیدروموتور|اوربیتال|پیستونی"),
      ("HYD_VALVE_HOSE", "هیدرولیک|فلوکنترل|فلو کنترل|پرشر|آکومولاتور|شیلنگ فشار|ربع گرد|فشار شکن"),
      ("PUMP_WELL", "شناور|کف کش|کفکش|لجن کش|کمرکش|موتور پمپ|موتورپمپ|گازوئیل کش"),
      ("PUMP_BUILDING", "پمپ|بوستر|سیرکولاتور|آبگرمکن|ست کنترل|منبع تحت فشار|تیوپ"),
      ("IRRIG", "قطره|آبیاری|نوار تیپ|آبپاش|رایزر|کمربند|ویسپار|مه پاش|لی فلت"),
      ("PIPE_PE_NETWORK", "پلی رود|پلی اتیلن|تکاب|الکتروفیوژن|دینا|جوشی|بست ابتدایی"),
      ("PIPE_SEWER", "upvc|پلیکا|پیمتاش|pvc"),
      ("PIPE_PLUMB", "گالوانیزه|زانو|سه راهی|سه راه|نیوپایپ|کپ رزوه|درپوش|زاکو|پنج لایه|سرشیلنگی|فیتینگ|تفلون|شیر حیاطی|آچار"),
      ("VALVE_IND", "شیر|گیج|درجه|لرزه گیر|صافی|اکچویتور"),
      ("ELEC_MOTOR", "الکتروموتور|گیربکس|بلبرینگ|کوپلینگ"),
      ("ELEC_DRIVE", "اینورتر|درایو|سافت"),
      ("ELEC_PANEL", "کلید|کنتاکتور|رله|فیوز|ترمینال|فانال|تابلو|کابل|ترانسمیتر|بوبین|اسکواردی|سکسیونر|برق"),
      ("FIRE_SAFETY", "عینک|دستکش|محافظ گوش|ایرپلاگ|آتش|کپسول|اسپرینکلر|ایمنی|آریا کوپلینگ"),
      ("LIGHT", "پروژکتور|چراغ|led|ال ای دی"),
      ("POOL", "کلر|استخر|آب نما|آبنما"),
      ("TOOLS", "برس|دریل|فرز|متر لیزری|سه پایه|گریس|فنرزن|ابزار"),
      ("ELEVATOR", "آسانسور|شستی"),
      ("LAB", "کیت|محلول|بافر"),
      ("IND_COOLER", "فن |بلور|هیتر|المنت|ترموستات"),
      ("SEED_MELON", "بذر")]
_KW = [(pid, re.compile(p, re.I)) for pid, p in KW]
# ستون‌هایی که به هیچ وجه خوانده نمی‌شوند
_PRIVATE = re.compile(r"نام|تلفن|شماره|تماس گیرنده|آدرس|بانک|کد ملی|ایمیل|برگه")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s if s is not None else "")).replace("ي", "ی").replace("ك", "ک").replace("‌", " ").strip()


def classify(text: str) -> str:
    t = _norm(text)
    for pid, rx in _KW:
        if rx.search(t):
            return pid
    return "OTHER"


def num(v) -> int:
    """«1.400.000»، «200,000»، «1) 3.500.000\\n2) 4.500.000»، «378.235\\n-14%»، «حدود 70 هزار» ← عدد تومان."""
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    tot = 0
    for line in str(v).translate(FA).split("\n"):
        line = re.sub(r"^\s*\d\)\s*", "", line)
        if "%" in line:
            continue
        m = re.search(r"\d[\d.,]*", line)
        if not m:
            continue
        s = m.group(0).strip(".,")
        digits = re.sub(r"\D", "", s)
        if "هزار" in line and len(digits) <= 3:
            tot += int(digits) * 1000
            continue
        parts = re.split(r"[.,]", s)
        if len(parts) > 1 and len(parts[-1]) == 2:   # «650.00» یعنی ۶۵۰ هزار
            parts[-1] += "0"
        tot += int("".join(parts) or 0)
    return tot


_CITY_IDX: dict | None = None


def _city_lookup(name: str) -> str | None:
    """شهرهای نقشه‌ی تبلیغات و پروفایل شهرها ← استان (مثلاً کوهدشت ← لرستان)."""
    global _CITY_IDX
    if _CITY_IDX is None:
        try:
            from types import SimpleNamespace
            from app.plan.cities import _norm as cnorm, city_index
            from app.plan.data import plan_file
            idx = city_index([SimpleNamespace(id=c["id"], segments=c.get("segments") or []) for c in plan_file()["cats"]])
            _CITY_IDX = {k: v["province"] for k, v in idx.items() if v.get("province")}
            _CITY_IDX["__norm__"] = cnorm
        except Exception as e:  # noqa: BLE001
            log.info("فهرست شهرها بارگذاری نشد: %s", e)
            _CITY_IDX = {}
    f = _CITY_IDX.get("__norm__")
    v = _CITY_IDX.get(f(name)) if f else None
    v = v if isinstance(v, str) else None
    return re.sub(r"\s*\(.*?\)", "", v).strip() if v else None


def province(p, city=None) -> str:
    for raw in (p, city):
        x = _norm(raw)
        if not x:
            continue
        if x in PROVS:
            return x
        for k, v in CITY2PROV.items():
            if x.startswith(k):
                return v
        for q in PROVS:
            if q in x:
                return q
        found = _city_lookup(x)
        if found in PROVS:
            return found
    return "نامشخص"


def _ym(v) -> str | None:
    if isinstance(v, (datetime, date)):
        try:
            import jdatetime
            j = jdatetime.date.fromgregorian(date=v if isinstance(v, date) and not isinstance(v, datetime) else v.date())
            return f"{j.year}/{j.month:02d}"
        except Exception:  # noqa: BLE001
            return None
    m = re.search(r"(14\d\d)\s*/\s*(\d\d?)", _norm(v).translate(FA))
    if not m or not (1 <= int(m.group(2)) <= 12):
        return None
    y = int(m.group(1))
    return f"{y}/{int(m.group(2)):02d}" if 1400 <= y <= 1406 else None


def _find(headers: list[str], *needles: str) -> int | None:
    for i, h in enumerate(headers):
        if _PRIVATE.search(h) and not any(n in h for n in ("واریزی", "قیمت")):
            continue
        if all(n in h for n in needles):
            return i
    return None


def _day(v) -> str | None:
    m = re.search(r"(14\d\d)\s*/\s*(\d\d?)\s*/\s*(\d\d?)", _norm(v).translate(FA))
    return f"{m.group(1)}/{int(m.group(2)):02d}/{int(m.group(3)):02d}" if m else None


def parse_rows(data: bytes) -> dict:
    """فایل ← ردیف‌های بی‌نام (بدون نام/تلفن/آدرس) به تفکیک نوع شیت."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    invoices, calls, petty, demand = [], [], [], []
    sheets = []
    for ws in wb.worksheets:
        sheets.append(ws.title)
        rows = ws.iter_rows(values_only=True)
        try:
            head = [_norm(h) for h in next(rows)]
        except StopIteration:
            continue
        ix = {"rev": _find(head, "مجموع", "واریزی مشتری"), "profit": _find(head, "مجموع", "سود"),
              "prov": _find(head, "استان"), "city": _find(head, "شهر"), "product": _find(head, "محصول"),
              "field": _find(head, "زمینه"), "paid_at": _find(head, "تاریخ واریزی"), "call_at": _find(head, "تاریخ تماس"),
              "ship": _find(head, "نحوه ارسال"), "want": _find(head, "محصولی که"), "buy": _find(head, "منجر"),
              "cost": _find(head, "هزینه"), "date": _find(head, "تاریخ"), "req": _find(head, "درخواست")}

        def g(r, k):
            i = ix.get(k)
            return r[i] if i is not None and i < len(r) else None

        if ix["rev"] is not None and ix["prov"] is not None:          # شیت «مجموع فاکتورها»
            for r in rows:
                rev = num(g(r, "rev"))
                if rev <= 0:
                    continue
                pr = num(g(r, "profit"))
                invoices.append({"ym": _ym(g(r, "paid_at")) or _ym(g(r, "call_at")),
                                 "prov": province(g(r, "prov"), g(r, "city")),
                                 "plan": classify(f"{_norm(g(r, 'product'))} {_norm(g(r, 'field'))}"),
                                 "rev": rev, "profit": pr if 0 < pr < rev else None,
                                 "ship": (_norm(g(r, "ship")).split(" ")[0] or None)})
        elif ix["buy"] is not None and ix["want"] is not None:        # شیت‌های تماس (هر تأمین‌کننده)
            for r in rows:
                want = _norm(g(r, "want"))
                if not want or want == "None":
                    continue
                calls.append({"sheet": ws.title, "ym": _ym(g(r, "call_at")),
                              "plan": (lambda c: c if c != "OTHER" else classify(ws.title))(classify(want)),
                              "buy": _norm(g(r, "buy")).startswith("بله")})
        elif ix["req"] is not None and ix["city"] is not None:        # دفتر تماس با شهر و درخواست (بدون نتیجه)
            for r in rows:
                d = g(r, "call_at") if ix["call_at"] is not None else g(r, "date")
                req, city = _norm(g(r, "req")), _norm(g(r, "city"))
                if not _day(d) and not req and not city:
                    continue
                demand.append({"day": _day(d), "ym": _ym(d), "prov": province(None, city) if city not in ("", "-", "None") else "نامشخص",
                               "plan": classify(req) if req and req != "None" else "OTHER"})
        elif ix["cost"] is not None:                                  # تنخواه
            for r in rows:
                c = num(g(r, "cost"))
                if c > 0:
                    petty.append({"ym": _ym(g(r, "date")), "amount": c})
    wb.close()
    import hashlib
    key = hashlib.sha1("|".join(sorted(sheets)).encode()).hexdigest()[:10]
    return {"key": key, "sheets": sheets, "invoices": invoices, "calls": calls, "petty": petty, "demand": demand}


def parse_workbook(data: bytes) -> dict:
    r = parse_rows(data)
    return _summarize(r["invoices"], r["calls"], r["petty"], r["demand"])


def _summarize(invoices: list[dict], calls: list[dict], petty: list[dict], demand: list[dict] | None = None) -> dict:
    demand = demand or []
    by_plan: dict[str, dict] = defaultdict(lambda: {"n": 0, "rev": 0, "pn": 0, "prev": 0, "profit": 0})
    plan_prov, by_prov, months, ship = Counter(), Counter(), defaultdict(lambda: {"n": 0, "rev": 0, "profit": 0}), Counter()
    revs = []
    for x in invoices:
        b = by_plan[x["plan"]]
        b["n"] += 1
        b["rev"] += x["rev"]
        revs.append(x["rev"])
        if x["profit"]:
            b["pn"] += 1
            b["prev"] += x["rev"]
            b["profit"] += x["profit"]
        by_prov[x["prov"]] += 1
        if x["prov"] != "نامشخص":
            plan_prov[f"{x['plan']}|{x['prov']}"] += 1
        if x["ym"]:
            mm = months[x["ym"]]
            mm["n"] += 1
            mm["rev"] += x["rev"]
            mm["profit"] += x["profit"] or 0
        if x["ship"]:
            ship[x["ship"]] += 1
    tot_prev = sum(b["prev"] for b in by_plan.values())
    tot_profit = sum(b["profit"] for b in by_plan.values())
    tot_pn = sum(b["pn"] for b in by_plan.values())
    margin = tot_profit / tot_prev if tot_prev else None
    avg_profit = tot_profit / tot_pn if tot_pn else None

    c_plan = defaultdict(lambda: {"calls": 0, "buys": 0})
    c_sheet = defaultdict(lambda: {"calls": 0, "buys": 0})
    c_month = defaultdict(lambda: {"calls": 0, "buys": 0})
    for c in calls:
        for agg, k in ((c_plan, c["plan"]), (c_sheet, c["sheet"]), (c_month, c["ym"])):
            if k:
                agg[k]["calls"] += 1
                agg[k]["buys"] += int(c["buy"])
    n_calls, n_buys = len(calls), sum(int(c["buy"]) for c in calls)
    cr = n_buys / n_calls if n_calls else None

    plans = {}
    for k, b in by_plan.items():
        m = b["profit"] / b["prev"] if b["prev"] else None
        ap = b["profit"] / b["pn"] if b["pn"] else None
        cp = c_plan.get(k) or {}
        pcr = cp["buys"] / cp["calls"] if cp.get("calls", 0) >= 20 else cr
        plans[k] = {"invoices": b["n"], "revenue": b["rev"], "profit": b["profit"],
                    "margin_pct": round(100 * m, 1) if m is not None else None,
                    "avg_order": round(b["rev"] / b["n"]), "avg_profit": round(ap) if ap else None,
                    "calls": cp.get("calls", 0), "buys": cp.get("buys", 0),
                    "call_conv_pct": round(100 * pcr, 1) if pcr else None,
                    "max_cost_per_call": round(ap * pcr) if ap and pcr else None}
    return {
        "imported_at": datetime.now(timezone.utc).isoformat(),
        "invoices": len(invoices), "revenue": sum(revs), "profit": tot_profit,
        "margin_pct": round(100 * margin, 1) if margin else None,
        "avg_order": round(sum(revs) / len(revs)) if revs else None,
        "median_order": round(statistics.median(revs)) if revs else None,
        "avg_profit": round(avg_profit) if avg_profit else None,
        "max_cost_per_call": round(avg_profit * cr) if avg_profit and cr else None,
        "plans": plans,
        "provinces": by_prov.most_common(31),
        "plan_prov": dict(plan_prov),
        "months": sorted(([k, v["n"], v["rev"], v["profit"]] for k, v in months.items()))[-18:],
        "shipping": ship.most_common(8),
        "calls": {"total": n_calls, "buys": n_buys, "conv_pct": round(100 * cr, 1) if cr else None,
                  "by_sheet": sorted(([k, v["calls"], v["buys"]] for k, v in c_sheet.items()), key=lambda x: -x[1])[:15],
                  "months": sorted([k, v["calls"], v["buys"]] for k, v in c_month.items())[-18:]},
        "demand": _demand(demand),
        "petty": {"rows": len(petty), "total": sum(p["amount"] for p in petty),
                  "per_invoice": round(sum(p["amount"] for p in petty) / len(invoices)) if invoices and petty else None},
    }


def _demand(rows: list[dict]) -> dict:
    """تماس‌های ورودی با شهر و درخواست: تقاضا به تفکیک استان و دسته + خط مبنای تماس روزانه."""
    if not rows:
        return {"total": 0}
    try:
        from app.plan.data import today_jalali
        t = today_jalali()
        today = f"{t.year}/{t.month:02d}/{t.day:02d}"
    except Exception:  # noqa: BLE001
        today = "9999"
    days = Counter(r["day"] for r in rows if r.get("day") and r["day"] <= today)   # تاریخ‌های اشتباهِ آینده حذف
    pp = Counter(f"{r['plan']}|{r['prov']}" for r in rows if r["prov"] != "نامشخص" and r["plan"] != "OTHER")
    return {"total": len(rows), "days": len(days),
            "per_day": round(sum(days.values()) / len(days), 1) if days else None,
            "first_day": min(days) if days else None, "last_day": max(days) if days else None,
            "by_day": sorted(days.items())[-60:],
            "provinces": Counter(r["prov"] for r in rows).most_common(20),
            "plans": Counter(r["plan"] for r in rows if r["plan"] != "OTHER").most_common(20),
            "plan_prov": dict(pp)}


SOURCES = "offline.sources"   # {کلید فایل: ردیف‌های بی‌نام}؛ هر فایل جدا نگه داشته و با بقیه جمع می‌شود


def import_file(data: bytes) -> dict:
    r = parse_rows(data)
    if not r["invoices"] and not r["calls"] and not r["demand"]:
        raise ValueError("شیت فاکتور یا تماس در فایل پیدا نشد (ستون‌های «مجموع واریزی مشتری» و «استان»، "
                         "یا «به خرید منجر شده»، یا «شهر» و «درخواست» لازم است)")
    src = get_setting(SOURCES) or {}
    src[r["key"]] = {"at": datetime.now(timezone.utc).isoformat(), "sheets": r["sheets"],
                     **{k: r[k] for k in ("invoices", "calls", "petty", "demand")}}
    set_setting(SOURCES, src)
    s = rebuild()
    log.info("فروش آفلاین: فایل %s (%d فاکتور، %d تماس، %d درخواست)", r["key"], len(r["invoices"]),
             len(r["calls"]), len(r["demand"]))
    return s


def rebuild() -> dict:
    src = get_setting(SOURCES) or {}
    allr = {k: [x for v in src.values() for x in v.get(k) or []] for k in ("invoices", "calls", "petty", "demand")}
    s = _summarize(allr["invoices"], allr["calls"], allr["petty"], allr["demand"])
    s["files"] = [{"key": k, "at": v.get("at"), "sheets": v.get("sheets"), "invoices": len(v.get("invoices") or []),
                   "calls": len(v.get("calls") or []), "demand": len(v.get("demand") or [])} for k, v in src.items()]
    set_setting(KEY, s)
    return s


def summary() -> dict | None:
    return get_setting(KEY)


def plan_stats(plan_id: str) -> dict | None:
    s = summary()
    return (s or {}).get("plans", {}).get(plan_id)


def plan_province_counts() -> dict[tuple[str, str], int]:
    s = summary() or {}
    out: dict = {}
    for k, v in (s.get("plan_prov") or {}).items():
        pl, _, pv = k.partition("|")
        out[(pl, pv)] = v
    for k, v in ((s.get("demand") or {}).get("plan_prov") or {}).items():   # هر تماس بدون فاکتور = نصف فاکتور
        pl, _, pv = k.partition("|")
        out[(pl, pv)] = out.get((pl, pv), 0) + 0.5 * v
    return out


def profit_factor(plan_id: str) -> float | None:
    """سود هر فاکتور این دسته نسبت به میانگین کل (با حداقل ۴ فاکتور)، نرم‌شده بین ۰٫۷۵ تا ۱٫۴."""
    s = summary() or {}
    p = (s.get("plans") or {}).get(plan_id) or {}
    if not s.get("avg_profit") or not p.get("avg_profit") or (p.get("invoices") or 0) < 4:
        return None
    return round(max(0.75, min(1.4, (p["avg_profit"] / s["avg_profit"]) ** 0.35)), 3)


def margin_rate(plan_id: str | None) -> float | None:
    """نرخ سود (۰ تا ۱) برای تخمین سود واقعی یک کمپین."""
    s = summary() or {}
    p = ((s.get("plans") or {}).get(plan_id) or {}) if plan_id else {}
    m = p.get("margin_pct") if (p.get("invoices") or 0) >= 4 else None
    m = m if m is not None else s.get("margin_pct")
    return m / 100 if m else None
