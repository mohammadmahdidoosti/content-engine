"""منطق CRM: مشتری، معامله، مرحله، پیگیری، گزارش، ورود فایل Leads."""
import hashlib
import hmac
import io
import os
import re
import secrets
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone

from sqlalchemy import func, or_

from crm.db import Activity, Contact, Deal, User, session

TEHRAN = timezone(timedelta(hours=3, minutes=30))
STAGES = {"new": "درخواست جدید", "quoted": "قیمت داده شد", "followup": "پیگیری", "invoice": "پیش‌فاکتور",
          "won": "خرید", "lost": "انصراف"}
OPEN = ("new", "quoted", "followup", "invoice")
LOST_REASONS = ["گران بود", "موجود نبود", "از جای دیگر خرید", "جواب نداد", "منصرف شد", "زمان تحویل", "سایر"]
SOURCES = ["کد آگهی", "ترب", "سایت", "تلگرام", "بله", "ایتا", "اینستاگرام", "تماس مستقیم", "معرفی", "مشتری قبلی", "سایر"]
_FA = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def today() -> date:
    return datetime.now(TEHRAN).date()


# ---------------- کاربر ----------------
def hash_pw(pw: str, salt: str | None = None) -> str:
    salt = salt or secrets.token_hex(8)
    h = hashlib.pbkdf2_hmac("sha256", pw.encode(), salt.encode(), 200_000).hex()
    return f"{salt}${h}"


def check_pw(pw: str, stored: str) -> bool:
    salt, _, h = (stored or "").partition("$")
    return bool(h) and hmac.compare_digest(hash_pw(pw, salt).split("$")[1], h)


def ensure_admin() -> None:
    """کاربر مدیر از متغیرهای CRM_ADMIN_USER / CRM_ADMIN_PASSWORD (فقط اگر هنوز کاربری نیست)."""
    u, p = os.environ.get("CRM_ADMIN_USER", "").strip(), os.environ.get("CRM_ADMIN_PASSWORD", "")
    if not u or len(p) < 8:
        return
    with session() as db:
        if db.query(User).count() == 0:
            db.add(User(username=u, name=u, pw_hash=hash_pw(p), role="admin"))
            db.commit()


def login(username: str, password: str) -> User | None:
    with session() as db:
        u = db.query(User).filter(func.lower(User.username) == (username or "").strip().lower()).first()
        return u if u and check_pw(password or "", u.pw_hash) else None


def add_user(username: str, password: str, name: str = "", role: str = "sales") -> dict:
    username = (username or "").strip()
    if not re.match(r"^[A-Za-z0-9_.]{3,40}$", username) or len(password or "") < 8:
        raise ValueError("نام کاربری (حروف انگلیسی) و رمز حداقل ۸ حرفی لازم است")
    with session() as db:
        if db.query(User).filter(func.lower(User.username) == username.lower()).first():
            raise ValueError("این نام کاربری هست")
        db.add(User(username=username, name=name or username, pw_hash=hash_pw(password), role=role if role in ("admin", "sales") else "sales"))
        db.commit()
    return {"ok": True}


def change_password(username: str, old: str, new: str) -> dict:
    if len(new or "") < 8:
        raise ValueError("رمز جدید حداقل ۸ حرف")
    with session() as db:
        u = db.query(User).filter(User.username == username).first()
        if not u or not check_pw(old, u.pw_hash):
            raise ValueError("رمز فعلی درست نیست")
        u.pw_hash = hash_pw(new)
        db.commit()
    return {"ok": True}


# ---------------- کمکی‌ها ----------------
def norm_phone(p) -> str | None:
    s = re.sub(r"\D", "", str(p or "").translate(_FA))
    if not s:
        return None
    if s.startswith("0098"):
        s = "0" + s[4:]
    elif s.startswith("98") and len(s) == 12:
        s = "0" + s[2:]
    elif s.startswith("9") and len(s) == 10:
        s = "0" + s
    return s if 8 <= len(s) <= 12 else None


def parse_day(v) -> date | None:
    """«۱۴۰۵/۰۸/۱۲» یا «2026-11-03» یا date ← date میلادی."""
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).translate(_FA).strip()
    m = re.match(r"^(\d{3,4})[/-](\d{1,2})[/-](\d{1,2})", s)
    if not m:
        return None
    y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if not (1300 <= y <= 1500 or 1900 <= y <= 2100):
        return None
    try:
        if y < 1700:
            import jdatetime
            return jdatetime.date(y, mo, d).togregorian()
        return date(y, mo, d)
    except ValueError:
        return None


def jalali(d) -> str | None:
    if not d:
        return None
    import jdatetime
    if isinstance(d, datetime):
        d = d.astimezone(TEHRAN).date()
    j = jdatetime.date.fromgregorian(date=d)
    return f"{j.year}/{j.month:02d}/{j.day:02d}"


def _money(v) -> int | None:
    """«۴۵,۰۰۰,۰۰۰» ← 45000000؛ اگر چند عدد در یک خانه است، اولی (سقف ۱۰۰۰ میلیارد)."""
    if isinstance(v, (int, float)):
        n = int(v)
    else:
        m = re.search(r"\d[\d,٬]*", str(v if v is not None else "").translate(_FA))
        if not m:
            return None
        n = int(re.sub(r"[,٬]", "", m.group(0)))
    return n if 0 < n < 10**12 else None


def _contact_row(c: Contact) -> dict:
    return {"id": c.id, "name": c.name, "phone": c.phone, "city": c.city, "province": c.province,
            "activity": c.activity, "tags": c.tags or [], "note": c.note}


def _deal_row(d: Deal, c: Contact | None = None) -> dict:
    t = today()
    return {"id": d.id, "contact_id": d.contact_id, "title": d.title, "category": d.category, "source": d.source,
            "ad_code": d.ad_code, "stage": d.stage, "stage_name": STAGES.get(d.stage, d.stage), "amount": d.amount,
            "profit": d.profit, "lost_reason": d.lost_reason, "note": d.note,
            "next_at": d.next_at.isoformat() if d.next_at else None, "next_at_fa": jalali(d.next_at),
            "overdue": bool(d.next_at and d.next_at < t and d.stage in OPEN),
            "created_fa": jalali(d.created_at), "updated_fa": jalali(d.updated_at),
            "contact": _contact_row(c) if c else None}


def _log(db, deal_id: int, kind: str, text: str, user: str | None) -> None:
    db.add(Activity(deal_id=deal_id, kind=kind, text=text[:2000], user=user))


# ---------------- مشتری و معامله ----------------
def find_contact(phone: str) -> dict | None:
    p = norm_phone(phone)
    if not p:
        return None
    with session() as db:
        c = db.query(Contact).filter(Contact.phone == p).first()
        if not c:
            return None
        deals = db.query(Deal).filter(Deal.contact_id == c.id).order_by(Deal.created_at.desc()).all()
        return {**_contact_row(c), "deals": [_deal_row(d) for d in deals]}


def quick_add(b: dict, user: str | None = None) -> dict:
    """ثبت سریع: شماره (لازم) + نام + کالا + منبع؛ مشتری قبلی با همان شماره پیدا می‌شود."""
    phone = norm_phone(b.get("phone"))
    if not phone:
        raise ValueError("شماره‌ی تلفن درست نیست")
    title = str(b.get("title") or "").strip()[:300]
    if not title:
        raise ValueError("کالای درخواستی را بنویسید")
    src = str(b.get("source") or "").strip()[:40] or None
    code = re.sub(r"[^A-Za-z0-9]", "", str(b.get("ad_code") or "")).upper()[:20] or None
    if code and not src:
        src = "کد آگهی"
    with session() as db:
        c = db.query(Contact).filter(Contact.phone == phone).first()
        if not c:
            c = Contact(phone=phone)
            db.add(c)
        for k in ("name", "city", "province", "activity"):
            v = str(b.get(k) or "").strip()
            if v:
                setattr(c, k, v[:160])
        db.flush()
        nxt = parse_day(b.get("next_at")) or (today() + timedelta(days=int(b.get("follow_days") or 0))
                                              if b.get("follow_days") else None)
        d = Deal(contact_id=c.id, title=title, category=(str(b.get("category") or "").strip()[:80] or None),
                 source=src, ad_code=code, stage=b.get("stage") if b.get("stage") in STAGES else "new",
                 amount=_money(b.get("amount")), note=(str(b.get("note") or "").strip()[:2000] or None),
                 next_at=nxt, origin="manual")
        db.add(d)
        db.flush()
        _log(db, d.id, "note", "ثبت درخواست" + (f": {d.note}" if d.note else ""), user)
        db.commit()
        return {"deal_id": d.id, "contact_id": c.id}


def update_deal(did: int, b: dict, user: str | None = None) -> dict:
    with session() as db:
        d = db.get(Deal, did)
        if not d:
            raise ValueError("معامله پیدا نشد")
        if "stage" in b and b["stage"] in STAGES and b["stage"] != d.stage:
            if b["stage"] == "lost" and not (b.get("lost_reason") or d.lost_reason):
                raise ValueError("دلیل انصراف را انتخاب کنید")
            _log(db, d.id, "stage", f"{STAGES[d.stage]} ← {STAGES[b['stage']]}", user)
            d.stage = b["stage"]
            if d.stage in ("won", "lost"):
                d.closed_at, d.next_at = datetime.now(timezone.utc), None
            else:
                d.closed_at = None
        if "lost_reason" in b:
            d.lost_reason = str(b["lost_reason"] or "").strip()[:60] or None
        for k in ("title", "category", "source", "note"):
            if k in b:
                setattr(d, k, str(b[k] or "").strip()[:300 if k != "note" else 2000] or (None if k != "title" else d.title))
        if "ad_code" in b:
            d.ad_code = re.sub(r"[^A-Za-z0-9]", "", str(b["ad_code"] or "")).upper()[:20] or None
        for k in ("amount", "profit"):
            if k in b:
                setattr(d, k, _money(b[k]))
        if "next_at" in b or "follow_days" in b:
            nxt = parse_day(b.get("next_at")) if b.get("next_at") else (
                today() + timedelta(days=int(b["follow_days"])) if str(b.get("follow_days") or "").isdigit() else None)
            d.next_at = nxt
            if nxt:
                _log(db, d.id, "followup", f"پیگیری: {jalali(nxt)}", user)
        if b.get("log"):
            _log(db, d.id, b.get("log_kind") if b.get("log_kind") in ("note", "call") else "note", str(b["log"]), user)
        db.commit()
    return {"ok": True}


def update_contact(cid: int, b: dict) -> dict:
    with session() as db:
        c = db.get(Contact, cid)
        if not c:
            raise ValueError("مشتری پیدا نشد")
        if "phone" in b:
            p = norm_phone(b["phone"])
            if not p:
                raise ValueError("شماره درست نیست")
            other = db.query(Contact).filter(Contact.phone == p, Contact.id != cid).first()
            if other:
                raise ValueError("این شماره برای مشتری دیگری ثبت است")
            c.phone = p
        for k in ("name", "city", "province", "activity", "note"):
            if k in b:
                setattr(c, k, str(b[k] or "").strip()[:2000 if k == "note" else 160] or None)
        if "tags" in b and isinstance(b["tags"], list):
            c.tags = [str(t).strip()[:40] for t in b["tags"] if str(t).strip()][:20]
        db.commit()
    return {"ok": True}


def deal_detail(did: int) -> dict:
    with session() as db:
        d = db.get(Deal, did)
        if not d:
            raise ValueError("معامله پیدا نشد")
        c = db.get(Contact, d.contact_id)
        acts = db.query(Activity).filter(Activity.deal_id == did).order_by(Activity.at.desc()).limit(100).all()
        others = db.query(Deal).filter(Deal.contact_id == c.id, Deal.id != did).order_by(Deal.created_at.desc()).limit(20).all()
        return {**_deal_row(d, c), "activities": [{"at_fa": jalali(a.at), "at": a.at.isoformat(), "kind": a.kind, "text": a.text,
                                                   "user": a.user} for a in acts],
                "other_deals": [_deal_row(o) for o in others]}


def delete_deal(did: int) -> dict:
    with session() as db:
        db.query(Activity).filter(Activity.deal_id == did).delete()
        db.query(Deal).filter(Deal.id == did).delete()
        db.commit()
    return {"ok": True}


# ---------------- فهرست‌ها ----------------
def today_view() -> dict:
    t = today()
    with session() as db:
        due = (db.query(Deal, Contact).join(Contact, Contact.id == Deal.contact_id)
               .filter(Deal.stage.in_(OPEN), Deal.next_at.isnot(None), Deal.next_at <= t)
               .order_by(Deal.next_at, Deal.id).limit(200).all())
        fresh = (db.query(Deal, Contact).join(Contact, Contact.id == Deal.contact_id)
                 .filter(Deal.stage.in_(OPEN), Deal.next_at.is_(None))
                 .order_by(Deal.created_at.desc()).limit(100).all())
        counts = dict(db.query(Deal.stage, func.count()).filter(Deal.origin.notin_(["leads_import"])).group_by(Deal.stage).all())
        since = datetime.now(timezone.utc) - timedelta(days=30)
        month = db.query(Deal).filter(Deal.created_at >= since, Deal.origin != "leads_import").all()
    return {"today": jalali(t), "due": [_deal_row(d, c) for d, c in due], "no_followup": [_deal_row(d, c) for d, c in fresh],
            "counts": {k: counts.get(k, 0) for k in STAGES},
            "month": {"deals": len(month), "won": sum(d.stage == "won" for d in month),
                      "revenue": sum(d.amount or 0 for d in month if d.stage == "won")}}


def pipeline(stage: str | None = None, q: str = "", limit: int = 200) -> list[dict]:
    with session() as db:
        qry = db.query(Deal, Contact).join(Contact, Contact.id == Deal.contact_id)
        if stage in STAGES:
            qry = qry.filter(Deal.stage == stage)
        else:
            qry = qry.filter(Deal.stage.in_(OPEN))
        q = (q or "").strip()
        if q:
            p = norm_phone(q)
            like = f"%{q}%"
            conds = [Contact.name.ilike(like), Deal.title.ilike(like), Contact.city.ilike(like), Deal.ad_code.ilike(like)]
            if p:
                conds.append(Contact.phone.like(f"%{p[-7:]}%"))
            qry = qry.filter(or_(*conds))
        rows = qry.order_by(Deal.updated_at.desc()).limit(limit).all()
        return [_deal_row(d, c) for d, c in rows]


def search_contacts(q: str = "", tag: str = "", category: str = "", months: list[int] | None = None, limit: int = 300) -> list[dict]:
    """برای «مشتری‌های فصلی»: کسانی که در این ماه‌ها (شمسی) از این دسته خریده‌اند."""
    import jdatetime
    with session() as db:
        qry = db.query(Contact)
        if q:
            like = f"%{q.strip()}%"
            qry = qry.filter(or_(Contact.name.ilike(like), Contact.city.ilike(like), Contact.phone.like(like)))
        rows = qry.order_by(Contact.created_at.desc()).limit(3000).all()
        won = defaultdict(list)
        for d in db.query(Deal).filter(Deal.stage == "won").all():
            won[d.contact_id].append(d)
        out = []
        for c in rows:
            if tag and tag not in (c.tags or []):
                continue
            ds = won.get(c.id, [])
            if category:
                ds = [d for d in ds if category in (d.category or "") or category in (d.title or "")]
                if not ds:
                    continue
            if months:
                ds = [d for d in ds if jdatetime.date.fromgregorian(date=(d.closed_at or d.created_at).astimezone(TEHRAN).date()).month in months]
                if not ds:
                    continue
            out.append({**_contact_row(c), "bought": len(ds), "spent": sum(d.amount or 0 for d in ds),
                        "last": jalali(max((d.closed_at or d.created_at) for d in ds)) if ds else None,
                        "items": list(dict.fromkeys(d.title for d in ds))[:5]})
            if len(out) >= limit:
                break
        return out


# ---------------- گزارش ----------------
def report(days: int = 90, include_import: bool = False) -> dict:
    since = datetime.now(timezone.utc) - timedelta(days=days)
    with session() as db:
        q = db.query(Deal).filter(Deal.created_at >= since)
        if not include_import:
            q = q.filter(Deal.origin != "leads_import")
        deals = q.all()

    def agg(key):
        g = defaultdict(lambda: Counter())
        for d in deals:
            k = key(d) or "نامشخص"
            g[k]["deals"] += 1
            g[k]["won"] += d.stage == "won"
            g[k]["lost"] += d.stage == "lost"
            g[k]["open"] += d.stage in OPEN
            g[k]["revenue"] += (d.amount or 0) if d.stage == "won" else 0
            g[k]["profit"] += (d.profit or 0) if d.stage == "won" else 0
        rows = []
        for k, v in g.items():
            closed = v["won"] + v["lost"]
            rows.append({"key": k, **v, "conv_pct": round(100 * v["won"] / v["deals"], 1) if v["deals"] else 0,
                         "win_of_closed_pct": round(100 * v["won"] / closed, 1) if closed else None})
        return sorted(rows, key=lambda r: -r["deals"])
    lost = Counter(d.lost_reason or "بدون دلیل" for d in deals if d.stage == "lost")
    return {"days": days, "total": len(deals), "won": sum(d.stage == "won" for d in deals),
            "revenue": sum(d.amount or 0 for d in deals if d.stage == "won"),
            "by_source": agg(lambda d: d.ad_code or d.source), "by_category": agg(lambda d: d.category),
            "lost_reasons": lost.most_common(), "stages": Counter(d.stage for d in deals)}


def stats_for_marketing(days: int = 180) -> dict:
    """فقط عدد (بدون نام و شماره) برای پنل بازاریابی: معامله/خرید/فروش به تفکیک کد آگهی و منبع."""
    since = datetime.now(timezone.utc) - timedelta(days=days)
    with session() as db:
        deals = db.query(Deal).filter(Deal.created_at >= since, Deal.origin != "leads_import").all()
    by_code, by_source = defaultdict(Counter), defaultdict(Counter)
    for d in deals:
        for g, k in ((by_code, d.ad_code), (by_source, d.source)):
            if not k:
                continue
            g[k]["deals"] += 1
            g[k]["won"] += d.stage == "won"
            g[k]["revenue"] += (d.amount or 0) if d.stage == "won" else 0
            g[k]["profit"] += (d.profit or 0) if d.stage == "won" else 0
    return {"days": days, "codes": {k: dict(v) for k, v in by_code.items()},
            "sources": {k: dict(v) for k, v in by_source.items()}}


# ---------------- ورود فایل Leads (تاریخچه) ----------------
def _cell(v) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def import_leads(data: bytes) -> dict:
    """شیت‌های تماس (هر تأمین‌کننده/دسته): تاریخ، نام، شماره، کالا، «چی گفته»، «به خرید منجر شده؟»، واریزی.
    شیت «مجموع فاکتورها»: استان/شهر/زمینه‌ی فعالیت مشتری را کامل می‌کند. تکراری وارد نمی‌شود."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    added = dup = skipped = enriched = 0
    with session() as db:
        seen = {k for (k,) in db.query(Deal.ext_key).filter(Deal.ext_key.isnot(None)).all() if k}
        opened = contacts_only = 0
        by_phone = {c.phone: c for c in db.query(Contact).all() if c.phone}
        by_name = defaultdict(list)
        for c in by_phone.values():
            if c.name:
                by_name[c.name.strip()].append(c)
        invoices_sheet = None
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                head = [_cell(h) for h in next(rows)]
            except StopIteration:
                continue

            def col(*needles):
                for i, h in enumerate(head):
                    if all(n in h for n in needles):
                        return i
                return None
            ix = {"date": col("تاریخ تماس"), "name": col("نام"), "phone": col("شماره"), "want": col("محصولی که"),
                  "said": col("چی گفته"), "buy": col("منجر"), "paid": col("مجموع", "واریزی مشتری")}
            if col("استان") is not None and col("شهر") is not None:
                invoices_sheet = ws.title
                ip, ic, ia, iname = col("استان"), col("شهر"), col("زمینه"), col("نام مشتری")
                for r in rows:
                    nm = _cell(r[iname]) if iname is not None and iname < len(r) else ""
                    for c in by_name.get(nm, []):
                        changed = False
                        for k, i in (("province", ip), ("city", ic), ("activity", ia)):
                            v = _cell(r[i]) if i is not None and i < len(r) else ""
                            if v and v != "None" and not getattr(c, k):
                                setattr(c, k, v[:120])
                                changed = True
                        enriched += changed
                continue
            names = [i for i, h in enumerate(head) if h.startswith("نام")]
            phones = [i for i, h in enumerate(head) if "شماره" in h]
            req, icity = col("درخواست"), col("شهر")
            if req is not None and ix["phone"] is not None and ix["want"] is None:
                # دفتر درخواست‌های اخیر (تاریخ، نام، شهر، شماره، درخواست): نتیجه ثبت نشده ← درخواست باز برای پیگیری
                for r in rows:
                    g = lambda i: r[i] if i is not None and i < len(r) else None  # noqa: E731
                    want, phone = _cell(g(req)), norm_phone(g(ix["phone"]))
                    if not want or want == "None" or not phone:
                        skipped += 1
                        continue
                    day = parse_day(g(ix["date"]))
                    key = hashlib.sha1(f"{ws.title}|{phone}|{want[:300]}|{day}".encode()).hexdigest()
                    if key in seen:
                        dup += 1
                        continue
                    seen.add(key)
                    c = by_phone.get(phone)
                    if not c:
                        c = Contact(phone=phone, name=_cell(g(ix["name"]))[:160], tags=["از دفتر درخواست‌ها"])
                        db.add(c)
                        db.flush()
                        by_phone[phone] = c
                    city = _cell(g(icity))
                    if city and city not in ("None", "-") and not c.city:
                        c.city = city[:80]
                    created = datetime.combine(day, datetime.min.time(), tzinfo=TEHRAN) if day else datetime.now(timezone.utc)
                    db.add(Deal(contact_id=c.id, title=want[:300], stage="new", origin="demand_import", ext_key=key,
                                note="از دفتر درخواست‌ها؛ نتیجه‌ی تماس ثبت نشده بود", created_at=created, updated_at=created))
                    opened += 1
                continue
            if ix["want"] is None and names and phones and len(names) == len(phones) and len(names) >= 2:
                # فهرست ساده‌ی «نام / شماره» (چند ستون جفتی): فقط مشتری ساخته/کامل می‌شود
                for r in rows:
                    for ni, pi in zip(names, phones):
                        phone = norm_phone(r[pi] if pi < len(r) else None)
                        name = _cell(r[ni] if ni < len(r) else None)
                        if not phone:
                            continue
                        c = by_phone.get(phone)
                        if not c:
                            c = Contact(phone=phone, name=("" if name == "None" else name)[:160], tags=[ws.title[:40]])
                            db.add(c)
                            db.flush()
                            by_phone[phone] = c
                            contacts_only += 1
                        elif not c.name and name and name != "None":
                            c.name = name[:160]
                continue
            if ix["want"] is None or ix["buy"] is None or ix["phone"] is None:
                continue
            for r in rows:
                g = lambda k: r[ix[k]] if ix[k] is not None and ix[k] < len(r) else None  # noqa: E731
                want = _cell(g("want"))
                phone = norm_phone(g("phone"))
                if not want or want == "None" or not phone:
                    skipped += 1
                    continue
                day = parse_day(g("date"))
                c = by_phone.get(phone)
                if not c:
                    c = Contact(phone=phone, name=_cell(g("name"))[:160], tags=["از فایل Leads"])
                    db.add(c)
                    db.flush()
                    by_phone[phone] = c
                    if c.name:
                        by_name[c.name].append(c)
                key = hashlib.sha1(f"{ws.title}|{phone}|{want[:300]}|{day}".encode()).hexdigest()
                if key in seen:
                    dup += 1
                    continue
                seen.add(key)
                bought = _cell(g("buy")).startswith("بله")
                created = datetime.combine(day, datetime.min.time(), tzinfo=TEHRAN) if day else datetime.now(timezone.utc)
                d = Deal(contact_id=c.id, title=want[:300], category=ws.title[:80], source=None,
                         stage="won" if bought else "lost", lost_reason=None if bought else "سابقه (فایل Leads)",
                         amount=_money(g("paid")) if bought else None, note=_cell(g("said"))[:2000] or None,
                         origin="leads_import", ext_key=key, created_at=created, updated_at=created, closed_at=created)
                db.add(d)
                added += 1
        db.commit()
    wb.close()
    return {"added": added, "opened": opened, "contacts_only": contacts_only, "duplicates": dup, "skipped": skipped,
            "enriched_contacts": enriched, "invoices_sheet": invoices_sheet}
